"""3개년 R&R 엑셀(본부별 시트, 상·하 2블록)과 3개년 보임도를 표준 CSV로 변환한다.

R&R 시트 구조
  상단 블록: [상위조직 | 조직명 | 미션 | 주요 R&R]              → 담당·본부급
  하단 블록: [상위조직 | 팀명  | 미션 | 주요 R&R | 세부 R&R]    → 팀급
  헤더 문구는 연도·시트마다 다르지만('담당', '부서명', '본부/실/담당'…)
  '미션' 열의 위치와 '세부 R&R' 유무만으로 판별할 수 있다.

보임도
  실명이 들어 있다. 성명 컬럼은 읽는 즉시 버리고 어디에도 쓰지 않는다.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import openpyxl
import pandas as pd

ITEM_SPLIT = re.compile(r"[\n○◦•]|(?:^|\s)[-–]\s")


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _items(text: str) -> list[str]:
    parts = ITEM_SPLIT.split(text or "")
    return [p.strip(" -–—·\t") for p in parts if len(p.strip(" -–—·\t")) > 3]


def parse_rr(path: Path, year: int) -> tuple[pd.DataFrame, pd.DataFrame]:
    """(담당급, 팀급) 두 테이블을 돌려준다."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    upper_rows, team_rows = [], []

    for ws in wb.worksheets:
        if ws.title == "표지":
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        sheet_hq = _clean(rows[0][0]) if rows else ws.title

        block = None            # None | "upper" | "team"
        parent = ""
        for raw in rows:
            cells = [_clean(v) for v in (raw + [None] * 6)[:6]]

            # 헤더 판정: 셋째 열이 '미션'
            if cells[2] == "미션":
                block = "team" if "세부" in cells[4] else "upper"
                parent = ""
                continue
            if block is None or not (cells[2] or cells[3]):
                continue

            # 병합셀 복원: 상위조직이 비어 있으면 직전 값을 잇는다
            if cells[0]:
                parent = cells[0]
            # 2026 CEO 직속 시트처럼 둘째 열이 비고 첫째 열에 조직명이 오는 경우
            name = cells[1] or cells[0]
            if not name:
                continue

            record = {"연도": year, "시트": ws.title, "본부(시트)": sheet_hq,
                      "상위조직": parent if cells[1] else "", "조직명": name,
                      "미션": cells[2], "주요RR": cells[3]}
            if block == "team":
                record["세부RR"] = cells[4]
                team_rows.append(record)
            else:
                upper_rows.append(record)
    wb.close()
    return pd.DataFrame(upper_rows), pd.DataFrame(team_rows)


def to_roles_long(teams: pd.DataFrame, name_to_code: dict[str, str]) -> pd.DataFrame:
    """파이프라인이 읽는 (팀코드, 팀명, 상위조직, 구분, 역할) 롱 포맷."""
    rows = []
    for row in teams.itertuples():
        code = name_to_code.get(row.조직명, re.sub(r"\s+", "", row.조직명))
        for kind, text in [("미션", row.미션), ("주요", row.주요RR), ("세부", getattr(row, "세부RR", ""))]:
            for item in _items(text):
                rows.append({"팀코드": code, "팀명": row.조직명, "상위조직": row.상위조직,
                             "구분": kind, "역할": item})
    return pd.DataFrame(rows)


def parse_leaders(path: Path) -> pd.DataFrame:
    """보임도 3개 블록. 성명은 읽지 않는다."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["요약"].iter_rows(values_only=True))
    wb.close()
    blocks = {2024: (1, 2, 4), 2025: (6, 7, 9), 2026: (11, 12, 14)}   # 조직, 사번, 직책 열
    out = []
    for year, (org_col, id_col, title_col) in blocks.items():
        for raw in rows[7:]:
            if raw[id_col] is None:
                continue
            out.append({"연도": year, "조직": _clean(raw[org_col]),
                        "사번": _clean(raw[id_col]), "직책": _clean(raw[title_col])})
    return pd.DataFrame(out)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--rr", nargs=3, metavar=("Y2024", "Y2025", "Y2026"), required=True)
    parser.add_argument("--leaders", required=True)
    parser.add_argument("--out", default="data")
    args = parser.parse_args()
    out = Path(args.out)
    out.mkdir(exist_ok=True)

    latest = pd.read_csv(out / "members_2026.csv", dtype=str)
    name_to_code = (latest.groupby(["팀명", "팀코드"]).size().reset_index(name="n")
                    .sort_values("n", ascending=False).drop_duplicates("팀명")
                    .set_index("팀명")["팀코드"].to_dict())

    print("R&R 변환:")
    for year, path in zip((2024, 2025, 2026), args.rr):
        upper, teams = parse_rr(Path(path), year)
        upper.to_csv(out / f"rr_upper_{year}.csv", index=False, encoding="utf-8-sig")
        teams.to_csv(out / f"rr_team_{year}.csv", index=False, encoding="utf-8-sig")
        print(f"  {year}: 담당·본부급 {len(upper):3}건 · 팀 {len(teams):3}건")
        if year == 2026:
            roles = to_roles_long(teams, name_to_code)
            roles.to_csv(out / "team_roles.csv", index=False, encoding="utf-8-sig")
            matched = len(set(roles["팀코드"]) & set(latest["팀코드"]))
            print(f"        → team_roles.csv 교체: 역할 {len(roles)}건 / 팀 {roles['팀코드'].nunique()}개 "
                  f"/ 명단 매칭 {matched}개")

    leaders = parse_leaders(Path(args.leaders))
    leaders.to_csv(out / "leaders_3y.csv", index=False, encoding="utf-8-sig")
    print(f"\n보임도: {leaders.groupby('연도').size().to_dict()}  (성명 제외)")


if __name__ == "__main__":
    main()
